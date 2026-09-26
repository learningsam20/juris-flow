"""Document parsing for PDF, DOCX, TXT, MD (PRD §5.2, §6.5)."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from pypdf import PdfReader

from app.config import get_settings
from app.core.errors import ValidationError

logger = logging.getLogger(__name__)


class ParseResult:
    def __init__(self, text: str, extracted: bool = True) -> None:
        self.text = text
        self.extracted = extracted


def parse_document(filename: str, content: bytes) -> ParseResult:
    ext = Path(filename).suffix.lower().lstrip(".")
    settings = get_settings()
    if ext not in settings.allowed_documents:
        raise ValidationError(
            f"Unsupported document type `.{ext}`. Allowed: {', '.join(settings.allowed_documents)}"
        )
    if ext == "txt":
        return ParseResult(content.decode("utf-8", errors="replace"))
    if ext == "md":
        return ParseResult(content.decode("utf-8", errors="replace"))
    if ext == "docx":
        return _parse_docx(content)
    if ext == "pdf":
        return _parse_pdf(content)
    raise ValidationError(f"Unsupported document type `.{ext}`")


def _ocr_pdf_with_tesseract(pdf: Any, limit: int) -> str:
    import shutil

    import pytesseract  # type: ignore[import-untyped]

    tesseract_bin = shutil.which("tesseract")
    if not tesseract_bin:
        for candidate in (
            "/opt/homebrew/bin/tesseract",
            "/usr/local/bin/tesseract",
            "/usr/bin/tesseract",
        ):
            if Path(candidate).exists():
                tesseract_bin = candidate
                break
    if tesseract_bin:
        pytesseract.pytesseract.tesseract_cmd = str(tesseract_bin)
    else:
        raise ValidationError(
            "Tesseract OCR binary not found on the system. "
            "Please install tesseract (e.g. `brew install tesseract`) to process scanned PDFs."
        )

    logger.info("Running Tesseract OCR on %d pages using %s", limit, tesseract_bin)
    extracted_pages: list[str] = []
    for i in range(limit):
        try:
            page = pdf[i]
            image = page.render(scale=2.0).to_pil()
            page_text = pytesseract.image_to_string(image).strip()
            if page_text:
                extracted_pages.append(page_text)
        except Exception as exc:
            logger.warning("Tesseract OCR failed on page %d: %s", i + 1, exc)
            continue
    return "\n\n".join(extracted_pages).strip()


def _ocr_pdf_with_ollama_vision(pdf: Any, limit: int, model: str, base_url: str) -> str:
    import base64
    from io import BytesIO

    import requests

    logger.info("Running Ollama Vision OCR on %d pages using model %s", limit, model)
    extracted_pages: list[str] = []
    for i in range(limit):
        try:
            page = pdf[i]
            image = page.render(scale=1.5).to_pil()
            buf = BytesIO()
            image.save(buf, format="JPEG", quality=80)
            img_b64 = base64.b64encode(buf.getvalue()).decode("utf-8")

            resp = requests.post(
                f"{base_url.rstrip('/')}/api/chat",
                json={
                    "model": model,
                    "messages": [
                        {
                            "role": "user",
                            "content": "Extract and transcribe all text from this scanned document page verbatim without any explanation or commentary.",
                            "images": [img_b64],
                        }
                    ],
                    "stream": False,
                    "options": {"temperature": 0.0},
                },
                timeout=60,
            )
            if resp.status_code == 200:
                txt = resp.json().get("message", {}).get("content", "").strip()
                if txt:
                    extracted_pages.append(txt)
            else:
                logger.warning(
                    "Ollama Vision OCR returned status %d on page %d: %s",
                    resp.status_code,
                    i + 1,
                    resp.text[:200],
                )
        except Exception as exc:
            logger.warning("Ollama Vision OCR failed on page %d: %s", i + 1, exc)
            continue
    return "\n\n".join(extracted_pages).strip()


def _ocr_pdf(content: bytes, max_pages: int = 50) -> str:
    import pypdfium2 as pdfium  # type: ignore[import-untyped]

    try:
        pdf = pdfium.PdfDocument(content)
    except Exception as exc:
        raise ValidationError(f"Invalid PDF or corrupt image data for OCR: {exc}")

    total_pages = len(pdf)
    limit = min(total_pages, max_pages)
    settings = get_settings()
    provider = getattr(settings, "ocr_provider", "tesseract").lower().strip()

    if provider in ("ollama_vision", "vllm", "vision"):
        model = getattr(settings, "ocr_vision_model", "openbmb/minicpm-v4.6:latest")
        base_url = getattr(settings, "ocr_vision_base_url", "http://localhost:11434")
        text = _ocr_pdf_with_ollama_vision(pdf, limit, model, base_url)
        if not text:
            logger.warning(
                "Ollama Vision OCR produced empty text or timed out; falling back to Tesseract OCR"
            )
            text = _ocr_pdf_with_tesseract(pdf, limit)
        return text
    return _ocr_pdf_with_tesseract(pdf, limit)


def _parse_pdf(content: bytes) -> ParseResult:
    from io import BytesIO

    try:
        reader = PdfReader(BytesIO(content))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n\n".join(pages).strip()
    except Exception:
        text = ""

    if not text:
        settings = get_settings()
        if settings.ocr_enabled:
            max_pages = getattr(settings, "ocr_max_pages", 50)
            text = _ocr_pdf(content, max_pages=max_pages)
            if not text:
                raise ValidationError(
                    "Scanned (image-only) PDF detected and OCR was executed, but no readable text could be recognized."
                )
            return ParseResult(text)
        raise ValidationError(
            "Scanned (image-only) PDF detected and OCR is not enabled. "
            "Enable JAIL_OCR_ENABLED and configure an OCR provider to process scanned PDFs."
        )
    return ParseResult(text)


def _parse_docx(content: bytes) -> ParseResult:
    from io import BytesIO

    import docx

    document = docx.Document(BytesIO(content))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return ParseResult("\n".join(paragraphs).strip())
