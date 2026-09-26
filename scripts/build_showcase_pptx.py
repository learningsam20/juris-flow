#!/usr/bin/env python3
"""Build JurisFlow showcase deck as PPTX for submission."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "JurisFlow-Showcase.pptx"

BG = RGBColor(0x07, 0x0B, 0x14)
PANEL = RGBColor(0x0E, 0x15, 0x24)
INK = RGBColor(0xEE, 0xF3, 0xFA)
MUTED = RGBColor(0x84, 0x94, 0xA8)
ACCENT = RGBColor(0x3D, 0x9C, 0xF0)
GREEN = RGBColor(0x2F, 0xD6, 0x7B)
AMBER = RGBColor(0xE8, 0xB8, 0x4A)
VIOLET = RGBColor(0x9B, 0x7B, 0xFF)
LINE = RGBColor(0x1A, 0x24, 0x38)


def set_run(run, *, size=14, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_text(slide, left, top, width, height, text, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, name=name)
    return box


def add_bullets(slide, left, top, width, height, items, *, size=15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        run = p.add_run()
        run.text = f"•  {item}"
        set_run(run, size=size, bold=True, color=INK)
    return box


def paint_bg(slide, prs):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    sp_tree = slide.shapes._spTree
    el = bg._element
    sp_tree.remove(el)
    sp_tree.insert(2, el)


def accent_bar(slide, left, top, width, color):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def panel(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = LINE
    return shape


def card(slide, left, top, width, height, title, items, accent, *, size=15):
    panel(slide, left, top, width, height)
    accent_bar(slide, left, top, width, accent)
    add_text(
        slide,
        left + Inches(0.22),
        top + Inches(0.16),
        width - Inches(0.4),
        Inches(0.3),
        title.upper(),
        size=11,
        bold=True,
        color=MUTED,
        name="Consolas",
    )
    add_bullets(
        slide,
        left + Inches(0.18),
        top + Inches(0.5),
        width - Inches(0.35),
        height - Inches(0.65),
        items,
        size=size,
    )


def footer(slide, prs, label):
    add_text(
        slide,
        Inches(0.5),
        prs.slide_height - Inches(0.42),
        Inches(8),
        Inches(0.28),
        "EDUCATIONAL · NOT LEGAL ADVICE",
        size=9,
        color=MUTED,
        name="Consolas",
    )
    add_text(
        slide,
        prs.slide_width - Inches(2.2),
        prs.slide_height - Inches(0.42),
        Inches(1.7),
        Inches(0.28),
        label,
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
        name="Consolas",
    )


def slide_header(slide, kicker, title):
    add_text(slide, Inches(0.5), Inches(0.25), Inches(6), Inches(0.28), "JURISFLOW · ENTERPRISE AI", size=10, color=ACCENT, name="Consolas")
    add_text(slide, Inches(0.5), Inches(0.52), Inches(4), Inches(0.28), kicker, size=11, color=ACCENT, name="Consolas")
    add_text(slide, Inches(0.5), Inches(0.82), Inches(12.3), Inches(0.5), title, size=26, bold=True, color=INK)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # —— Title ——
    st = prs.slides.add_slide(blank)
    paint_bg(st, prs)
    add_text(st, Inches(0.7), Inches(2.15), Inches(12), Inches(0.35), "ENTERPRISE AI · FINAL PRESENTATION", size=14, color=ACCENT, name="Consolas")
    add_text(st, Inches(0.7), Inches(2.65), Inches(12), Inches(0.9), "JurisFlow", size=54, bold=True, color=INK)
    add_text(
        st,
        Inches(0.7),
        Inches(3.7),
        Inches(11.5),
        Inches(0.7),
        "Grounded legal KM · multi-agent tribunal · training artefacts",
        size=20,
        color=MUTED,
    )
    add_text(st, Inches(0.7), Inches(6.55), Inches(11), Inches(0.3), "Educational product — not legal advice", size=12, color=MUTED, name="Consolas")

    # —— 01 Problem ——
    s1 = prs.slides.add_slide(blank)
    paint_bg(s1, prs)
    slide_header(s1, "01 · PROBLEM", "Scale judgment. Don't burn SMEs on repetition.")
    gap, margin = Inches(0.22), Inches(0.5)
    usable_w = prs.slide_width - 2 * margin
    usable_h = Inches(4.9)
    top0 = Inches(1.55)
    cw = (usable_w - gap) / 2
    ch = (usable_h - gap) / 2
    quads = [
        ("Need", ["Scale", "Efficiency", "Policy currency"], ACCENT),
        ("Challenges", ["SME bandwidth", "Fast onboarding", "Agentified HITL CLM flow"], AMBER),
        ("Why enterprise", ["Proprietary knowledge", "High-stakes decisions", "Data privacy"], GREEN),
        (
            "What we built",
            [
                "Enterprise knowledge hub",
                "Grounded review",
                "Fleet of agents",
                "Tribunal simulation",
                "Training pack",
                "Risk insights",
            ],
            VIOLET,
        ),
    ]
    for i, (title, items, accent) in enumerate(quads):
        col, row = i % 2, i // 2
        card(s1, margin + col * (cw + gap), top0 + row * (ch + gap), cw, ch, title, items, accent, size=16)
    footer(s1, prs, "01 / 03")

    # —— 02 Architecture ——
    s2 = prs.slides.add_slide(blank)
    paint_bg(s2, prs)
    slide_header(s2, "02 · ARCHITECTURE", "Agents · policy gate · right model for the job")

    # Left column: stack
    card(
        s2,
        Inches(0.5),
        Inches(1.5),
        Inches(4.1),
        Inches(2.55),
        "Control & workloads",
        [
            "React UI → FastAPI → OPA/Rego",
            "Knowledge Hub · Grounded Review",
            "LangGraph Sim · edge-tts",
            "MCP tools · default-deny",
        ],
        ACCENT,
        size=14,
    )
    card(
        s2,
        Inches(0.5),
        Inches(4.2),
        Inches(4.1),
        Inches(2.55),
        "Agent fleet",
        [
            "Orchestrator",
            "Policy Reviewer",
            "Plaintiff · Defendant",
            "Judge · Informer · Witness",
        ],
        VIOLET,
        size=14,
    )

    # Middle: LLM routing
    card(
        s2,
        Inches(4.8),
        Inches(1.5),
        Inches(3.9),
        Inches(5.25),
        "LLM routing",
        [
            "Ollama — private / offline / cost",
            "OpenRouter — stronger models per task",
            "Vertex — optional GCP",
            "Per use-case adapter switch",
            "Qdrant — filtered org vectors",
            "edge-tts — overview + hearing",
            "No silent LLM fallback",
        ],
        AMBER,
        size=14,
    )

    # Right: tech bets
    panel(s2, Inches(8.9), Inches(1.5), Inches(3.95), Inches(5.25))
    accent_bar(s2, Inches(8.9), Inches(1.5), Inches(3.95), GREEN)
    add_text(s2, Inches(9.1), Inches(1.68), Inches(3.5), Inches(0.3), "TECH BETS", size=11, bold=True, color=MUTED, name="Consolas")
    bets = [
        ("Rationale", "Choice", "Not selected"),
        ("No lock-in", "Ollama + OpenRouter", "Single-vendor LLM"),
        ("HITL agent graph", "LangGraph", "Ad-hoc chains"),
        ("Org vectors", "Qdrant", "Chroma"),
        ("Default-deny", "OPA / Rego", "App-only RBAC"),
        ("Multi-voice TTS", "edge-tts", "Cloud TTS lock-in"),
        ("Ship FOSS fast", "FastAPI · React", "Heavy CLM suite"),
    ]
    y = Inches(2.1)
    for i, (a, b, c) in enumerate(bets):
        color = MUTED if i == 0 else INK
        bold = True
        add_text(s2, Inches(9.1), y, Inches(3.55), Inches(0.22), a, size=10, bold=bold, color=ACCENT if i else MUTED, name="Consolas")
        add_text(s2, Inches(9.1), y + Inches(0.2), Inches(3.55), Inches(0.22), b, size=12, bold=True, color=color)
        if i:
            add_text(s2, Inches(9.1), y + Inches(0.4), Inches(3.55), Inches(0.2), f"not: {c}", size=10, color=MUTED)
        y += Inches(0.7) if i == 0 else Inches(0.72)
    footer(s2, prs, "02 / 03")

    # —— 03 Roadmap ——
    s3 = prs.slides.add_slide(blank)
    paint_bg(s3, prs)
    slide_header(s3, "03 · ROADMAP", "Shipped core → gated enterprise CLM")

    track = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.7), Inches(10.9), Pt(3))
    track.fill.solid()
    track.fill.fore_color.rgb = LINE
    track.line.fill.background()
    done = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.7), Inches(3.2), Pt(3))
    done.fill.solid()
    done.fill.fore_color.rgb = GREEN
    done.line.fill.background()
    for x, label, color in [(Inches(2.2), "NOW", GREEN), (Inches(6.4), "NEAR", AMBER), (Inches(10.6), "FUTURE", VIOLET)]:
        dot = s3.shapes.add_shape(MSO_SHAPE.OVAL, x, Inches(1.57), Inches(0.28), Inches(0.28))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()
        add_text(s3, x - Inches(0.45), Inches(1.95), Inches(1.2), Inches(0.28), label, size=11, bold=True, color=color, align=PP_ALIGN.CENTER, name="Consolas")

    lanes = [
        (
            Inches(0.5),
            "Now · shipped",
            "Rehearsal core",
            [
                "Enterprise knowledge hub",
                "Grounded review",
                "Fleet of agents",
                "Tribunal simulation",
                "Training pack + audio",
                "Risk insights",
            ],
            GREEN,
        ),
        (
            Inches(4.7),
            "Near",
            "Trust loop",
            [
                "Repair AI reviews with HITL",
                "Auto triage / routing",
                "AI watermarking",
                "DeepEvals",
            ],
            AMBER,
        ),
        (
            Inches(8.9),
            "Future",
            "Legal OS",
            [
                "Complete CLM + CRM integration",
                "Hard AI gating",
                "SME maturity assessment",
            ],
            VIOLET,
        ),
    ]
    for left, badge, subtitle, items, accent in lanes:
        panel(s3, left, Inches(2.4), Inches(3.9), Inches(4.3))
        accent_bar(s3, left, Inches(2.4), Inches(3.9), accent)
        add_text(s3, left + Inches(0.22), Inches(2.58), Inches(3.5), Inches(0.28), badge.upper(), size=11, bold=True, color=accent, name="Consolas")
        add_text(s3, left + Inches(0.22), Inches(2.9), Inches(3.5), Inches(0.35), subtitle, size=18, bold=True, color=INK)
        add_bullets(s3, left + Inches(0.18), Inches(3.4), Inches(3.5), Inches(3.1), items, size=14)
    footer(s3, prs, "03 / 03")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size:,} bytes · slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
